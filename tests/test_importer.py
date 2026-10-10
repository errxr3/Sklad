"""Загрузка справочника артикулов из файла с тремя листами."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.importer import import_articles
from app.models import Article


@pytest.fixture
def top_xlsx(tmp_path) -> Path:
    """Файл, похожий на TOP.xlsx: три листа, артикулы в столбце A."""
    workbook = Workbook()
    first = workbook.active
    first.title = "Зеркала"
    rows_first = [
        ["Артикул"],            # заголовок — пропускаем
        ["Zerkalo 600x800"],
        ["  Zerkalo   400x600  "],   # лишние пробелы
        ["Zerkalo 500х700"],         # русская «х»
        [None],                      # пустая строка
        ["Zerkalo 600X800"],         # тот же, что выше, другим регистром
    ]
    for row in rows_first:
        first.append(row)

    second = workbook.create_sheet("Alum")
    for row in [["Артикул"], ["Alum 400x1200"], ["Alum 600x900"]]:
        second.append(row)

    third = workbook.create_sheet("Ambient")
    for row in [["Ambient 700x500"], ["Zerkalo 600x800"]]:  # повтор с первого листа
        third.append(row)

    path = tmp_path / "TOP.xlsx"
    workbook.save(path)
    return path


def test_импорт_трёх_листов(db, top_xlsx):
    report = import_articles(db, top_xlsx)

    codes = sorted(code for (code,) in db.query(Article.code).all())
    assert codes == [
        "Alum 400x1200",
        "Alum 600x900",
        "Ambient 700x500",
        "Zerkalo 400x600",
        "Zerkalo 500x700",
        "Zerkalo 600x800",
    ]
    assert report.added == 6
    assert set(report.per_sheet) == {"Зеркала", "Alum", "Ambient"}


def test_заголовки_и_пустые_строки_пропускаются(db, top_xlsx):
    report = import_articles(db, top_xlsx)
    assert report.skipped == 3  # два заголовка «Артикул» и одна пустая строка


def test_русская_х_не_плодит_дубли(db, top_xlsx):
    import_articles(db, top_xlsx)
    codes = [code for (code,) in db.query(Article.code).all()]
    assert "Zerkalo 500x700" in codes
    assert "Zerkalo 500х700" not in codes


def test_буквы_группы_проставляются_при_импорте(db, top_xlsx):
    import_articles(db, top_xlsx)
    letters = {
        code: group
        for code, group in db.query(Article.code, Article.group_letters).all()
    }
    assert letters["Zerkalo 600x800"] == "Z"
    assert letters["Alum 400x1200"] == "A"
    assert letters["Ambient 700x500"] == "AT"


def test_повторный_импорт_ничего_не_дублирует(db, top_xlsx):
    first = import_articles(db, top_xlsx)
    before = db.query(Article).count()

    second = import_articles(db, top_xlsx)
    assert db.query(Article).count() == before
    assert second.added == 0
    # «Уже были» считается по строкам файла: восемь значимых строк,
    # из которых два повтора — внутри самого файла.
    assert second.already == 8
    assert first.added + first.already == second.already


def test_импорт_из_csv(db, tmp_path):
    path = tmp_path / "spisok.csv"
    path.write_text("Артикул\nZerkalo 600х800\nAlum 500\n", encoding="utf-8")
    report = import_articles(db, path)
    assert report.added == 2
    assert sorted(c for (c,) in db.query(Article.code).all()) == [
        "Alum 500",
        "Zerkalo 600x800",
    ]
